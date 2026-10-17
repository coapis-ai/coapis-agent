"""B3 wiring: ``doc_reader`` reads through the pagination base.

Acceptance S4-4 is the one with teeth here — the per-read cap is the
governance ledger's number, not a constant baked into the tool.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from coapis.agents.tools.doc_reader import (
    _PAGINATE_SIZE_BYTES,
    _read_chars_budget,
    doc_reader,
)
from coapis.config.config import ToolResultPruningConfig
from coapis.config.context import set_current_workspace_dir
from coapis.governance.limits import LimitKind, LimitLedger


@pytest.fixture(autouse=True)
def _workspace(tmp_path, request):
    # Keep the parse cache inside the test's own tmp_path.
    set_current_workspace_dir(tmp_path)
    yield
    set_current_workspace_dir(None)


def _write_docx(path: Path, chapters: int = 12) -> Path:
    from docx import Document

    doc = Document()
    for chapter in range(1, chapters + 1):
        doc.add_heading(f"第{chapter}章 章节标题", level=1)
        doc.add_paragraph("正文内容 " * 300)
        doc.add_paragraph(f"needle-{chapter} 关键结论 " + "补充 " * 100)
    doc.save(str(path))
    return path


def test_default_cap_is_the_ledger_number(monkeypatch) -> None:
    """The cap comes from governance, not from a constant in the tool."""
    from coapis.governance.limits import LimitLedger

    ledger = LimitLedger.from_config(
        pruning_config=ToolResultPruningConfig(),
    )
    budget = ledger.limit(LimitKind.READ_CHARS).budget
    assert budget == 12000
    monkeypatch.setattr(
        "coapis.governance.limits.ledger_for_agent",
        lambda *a, **k: ledger,
    )
    assert _read_chars_budget() == budget
    # The old constant was 100000; the ledger is an order of magnitude
    # tighter, which is the whole point of the wiring.
    assert _read_chars_budget() < 100_000


def test_cap_falls_back_when_governance_is_off() -> None:
    # No agent context → the ledger is empty and raises on lookup; a read
    # must still have a sane cap.
    assert _read_chars_budget() == 20000


async def test_modes_delegate_to_the_pagination_base(tmp_path) -> None:
    doc = _write_docx(tmp_path / "wire.docx")
    index = await doc_reader(action="read", file_path=str(doc), mode="index")
    assert index["mode"] == "index"
    assert index["block_count"] >= 12
    first = index["blocks"][0]

    rng = await doc_reader(
        action="read",
        file_path=str(doc),
        mode="range",
        start=first["i"],
        end=first["i"],
    )
    assert rng["mode"] == "range"
    assert rng["unit"] == "block"
    assert rng["chars"] == first["chars"]
    assert rng["cursor"]["next_block"] == first["i"] + 1

    chars = await doc_reader(
        action="read",
        file_path=str(doc),
        mode="range",
        start_char=first["start"],
        end_char=first["end"],
    )
    assert chars["unit"] == "char"
    assert chars["chars"] == first["chars"]
    assert chars["cursor"]["next_start"] == first["end"]

    hit = await doc_reader(
        action="read",
        file_path=str(doc),
        mode="search",
        keywords="needle-5",
    )
    assert hit["mode"] == "search"
    assert hit["hits"] >= 1
    assert "needle-5" in hit["results"][0]["text"]


async def test_full_read_is_capped_by_the_ledger(tmp_path) -> None:
    doc = _write_docx(tmp_path / "cap.docx", chapters=40)
    out = await doc_reader(action="read", file_path=str(doc), mode="full")
    assert out["format"] == "python-docx"
    assert len(out["content"]) <= _read_chars_budget()
    assert out["truncated"] is True
    assert out["cursor"]["next_start"] == len(out["content"])
    assert out["cursor"]["remaining_chars"] > 0


async def test_oversized_document_is_not_read_whole(tmp_path) -> None:
    huge = tmp_path / "huge.txt"
    huge.write_text("行内容 needle-needle " * 20000, encoding="utf-8")
    assert huge.stat().st_size >= _PAGINATE_SIZE_BYTES
    out = await doc_reader(action="read", file_path=str(huge), mode="full")
    assert out["mode"] == "index+head"
    assert len(out["content"]) <= _read_chars_budget()
    assert out["index"]
    assert "不整读" in out["warning"]
    # The head is a block read, so its cursor continues in blocks.
    assert out["cursor"]["next_block"] > 0


async def test_bad_mode_is_rejected(tmp_path) -> None:
    doc = _write_docx(tmp_path / "m.docx", chapters=2)
    out = await doc_reader(action="read", file_path=str(doc), mode="vibe")
    assert "mode 无效" in out["error"]


async def test_cache_lives_in_the_workspace(tmp_path) -> None:
    doc = _write_docx(tmp_path / "cache.docx", chapters=3)
    await doc_reader(action="read", file_path=str(doc), mode="index")
    cached = list((tmp_path / ".doccache").glob("*.md"))
    assert len(cached) == 1, "one canonical markdown, inside the workspace"
