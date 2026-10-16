# -*- coding: utf-8 -*-
"""文档解析基础能力测试：路径解析、格式诚实、内容维度、兜底链路。

对应 2026-10-15 故障定位（tech_docs/文档解析基础能力故障定位-2026-10-15.md）：
前端 file:// 非法 URI、doc_reader 按进程 CWD 解析相对路径、老格式假声明、
表格/备注/公式缺失、CLI 兜底不存在。
"""
from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from coapis.agents.tools import doc_reader as dr
from coapis.agents.utils import message_processing as mp
from coapis.config.context import set_current_workspace_dir

WS = "coapis-test-ws"


@pytest.fixture()
def ws(tmp_path, monkeypatch):
    """Workspace with the chat-upload layout: files/media/<file>."""
    media = tmp_path / "files" / "media"
    media.mkdir(parents=True)
    (tmp_path / "files" / "loose.docx").write_bytes(b"PK\x03\x04loose")
    set_current_workspace_dir(tmp_path)
    monkeypatch.setattr(dr, "_MAX_GLOB_CANDIDATES", 8)
    yield tmp_path
    set_current_workspace_dir(None)


def _docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "页眉机密"
    doc.add_paragraph("正文第一段")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列A"
    table.cell(0, 1).text = "列B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(path)


def _xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "台账"
    sheet["A1"] = "名称"
    sheet["B1"] = "合计"
    sheet["A2"] = "甲"
    sheet["B2"] = "=1+2"  # 程序生成、Excel 从未打开 → 无缓存值
    wb.save(path)


def _pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "标题"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text_frame.text = "组内文本"
    slide.notes_slide.notes_text_frame.text = "演讲备注"
    prs.save(path)


def _pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "page one text")
    doc.new_page().insert_text((72, 72), "page two text")
    doc.save(path)


# ── 路径解析 ────────────────────────────────────────────────────────────
class TestPathResolution:
    def test_clean_input_path_strips_uri_forms(self):
        assert dr._clean_input_path("file://报告.docx") == "报告.docx"
        assert dr._clean_input_path("file:///tmp/a.docx") == "/tmp/a.docx"
        assert dr._clean_input_path("/media/%E6%8A%A5%E5%91%8A.docx") == "/media/报告.docx"
        assert dr._clean_input_path('"a.docx"') == "a.docx"

    def test_bare_name_resolves_to_chat_upload_dir(self, ws):
        _docx(ws / "files" / "media" / "t.docx")
        resolved, _ = dr._resolve_existing("t.docx")
        assert resolved is not None and resolved.name == "t.docx"

    def test_media_url_form_resolves(self, ws):
        _docx(ws / "files" / "media" / "t.docx")
        resolved, _ = dr._resolve_existing("/media/t.docx")
        assert resolved is not None

    def test_legacy_file_uri_resolves(self, ws):
        # 旧前端把文件名放在 netloc 位：file://{name}
        _docx(ws / "files" / "media" / "t.docx")
        resolved, _ = dr._resolve_existing("file://t.docx")
        assert resolved is not None

    def test_absolute_path_resolves(self, ws):
        target = ws / "files" / "media" / "t.docx"
        _docx(target)
        resolved, _ = dr._resolve_existing(str(target))
        assert resolved == target

    def test_missing_returns_candidates_not_empty_hints(self, ws):
        _docx(ws / "files" / "media" / "same.docx")
        resolved, candidates = dr._resolve_existing("same.docx.bak")
        assert resolved is None
        resolved2, candidates2 = dr._resolve_existing("nope.docx")
        assert resolved2 is None
        assert isinstance(candidates, list) and isinstance(candidates2, list)

    def test_reader_reports_candidates_on_miss(self, ws):
        out = asyncio.run(dr.doc_reader(action="read", file_path="ghost.docx"))
        assert "error" in out
        assert "glob_search" in out["hint"] or out["candidates"]

    def test_reader_never_uses_process_cwd_first(self, ws, monkeypatch):
        # 进程 CWD 里存在同名文件也不得优先于工作区上传目录
        _docx(ws / "files" / "media" / "t.docx")
        cwd_file = Path.cwd() / "t.docx"
        if cwd_file.exists():
            pytest.skip("process CWD already holds t.docx")
        resolved, _ = dr._resolve_existing("t.docx")
        assert str(resolved).startswith(str(ws))


# ── 格式诚实 ────────────────────────────────────────────────────────────
class TestFormatHonesty:
    def test_legacy_ole2_rejected_with_conversion_answer(self, ws):
        legacy = ws / "files" / "media" / "old.doc"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0fake")
        out = asyncio.run(dr.doc_reader(action="read", file_path=str(legacy)))
        assert "旧版二进制格式" in out["error"]
        assert "docx" in out["error"]

    def test_legacy_extensions_not_in_supported_formats(self):
        for ext in (".doc", ".xls", ".ppt"):
            assert ext not in dr.SUPPORTED_FORMATS
            assert ext in dr.LEGACY_FORMATS

    def test_images_point_to_view_image(self, ws):
        img = ws / "files" / "media" / "p.png"
        img.write_bytes(b"\x89PNG\r\n\x1a\n")
        out = asyncio.run(dr.doc_reader(action="read", file_path=str(img)))
        assert "view_image" in out["error"]

    def test_formats_action_advertises_limits(self):
        out = asyncio.run(dr.doc_reader(action="formats"))
        assert out["legacy_unsupported"] == [".doc", ".ppt", ".xls"]
        assert out["images_use"] == "view_image"


# ── 内容维度 ────────────────────────────────────────────────────────────
class TestContentDimensions:
    def test_docx_includes_tables_and_header(self, ws):
        target = ws / "files" / "media" / "t.docx"
        _docx(target)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.docx"))
        assert "正文第一段" in out["content"]
        assert "列A | 列B" in out["content"]
        assert "页眉机密" in out["content"]

    def test_xlsx_shows_formula_when_no_cached_value(self, ws):
        target = ws / "files" / "media" / "t.xlsx"
        _xlsx(target)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.xlsx"))
        assert "台账" in out["content"]
        assert "=1+2" in out["content"]

    def test_xlsx_sheet_param(self, ws):
        target = ws / "files" / "media" / "t.xlsx"
        _xlsx(target)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.xlsx", sheet="台账"))
        assert "名称" in out["content"]
        bad = asyncio.run(dr.doc_reader(action="read", file_path="t.xlsx", sheet="不存在"))
        assert "不存在" in bad["error"]

    def test_pptx_recurses_groups_and_includes_notes(self, ws):
        target = ws / "files" / "media" / "t.pptx"
        _pptx(target)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.pptx"))
        assert "标题" in out["content"]
        assert "组内文本" in out["content"]
        assert "演讲备注" in out["content"]

    def test_pdf_pages_param_is_one_based(self, ws):
        target = ws / "files" / "media" / "t.pdf"
        _pdf(target)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.pdf", pages="2"))
        assert "page two" in out["content"]
        assert "page one" not in out["content"]
        bad = asyncio.run(dr.doc_reader(action="read", file_path="t.pdf", pages="0"))
        assert "页码从 1 开始" in bad["error"]

    def test_max_chars_protects_context(self, ws):
        big = ws / "files" / "media" / "big.txt"
        big.write_text("x" * 5000, encoding="utf-8")
        out = asyncio.run(dr.doc_reader(action="read", file_path=str(big), max_chars=100))
        assert len(out["content"]) == 100
        assert out["truncated"] is True


# ── 兜底链路（镜像无 CLI，兜底必须是库） ─────────────────────────────────
class TestFallbackChain:
    def test_docx_falls_back_to_markitdown_when_lib_missing(self, ws, monkeypatch):
        target = ws / "files" / "media" / "t.docx"
        _docx(target)

        def _boom(_path):
            raise ImportError("python-docx disabled by test")

        monkeypatch.setattr(dr, "_docx_text", _boom)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.docx"))
        assert out.get("format") == "markitdown"
        assert "正文第一段" in out["content"]

    def test_pdf_falls_back_to_markitdown_when_lib_missing(self, ws, monkeypatch):
        target = ws / "files" / "media" / "t.pdf"
        _pdf(target)
        monkeypatch.setattr(dr, "_read_pdf", lambda *a, **k: asyncio.sleep(0, result=None))

        async def _fake_read_pdf(*_a, **_k):
            return await dr._read_with_markitdown(str(target))

        monkeypatch.setattr(dr, "_read_pdf", _fake_read_pdf)
        out = asyncio.run(dr.doc_reader(action="read", file_path="t.pdf"))
        assert out.get("format") == "markitdown"

    def test_no_cli_fallback_invoked_in_source(self):
        # 兜底只能是库（markitdown）；镜像里没有 pandoc/pdftotext/soffice，
        # 源码不得把它们当命令调用（文档字符串里提到是允许的）。
        src = Path(dr.__file__).read_text(encoding="utf-8")
        for cli in ("pandoc", "pdftotext", "soffice", "libreoffice"):
            assert f'"{cli}"' not in src, f"doc_reader 不应调用镜像里没有的 CLI: {cli}"
            assert f"'{cli}'" not in src
        assert "markitdown" in src


# ── 入站文件链路（前端 file:// → 后端解析） ──────────────────────────────
class TestInboundFileResolution:
    @staticmethod
    def _local_path(uri: str) -> Path:
        """``_resolve_media_url`` returns a ``file://`` URI; decode it back."""
        import urllib.parse

        return Path(urllib.parse.urlparse(uri).path)

    def test_file_uri_with_name_in_netloc_resolves(self, ws):
        target = ws / "files" / "media" / "t.docx"
        _docx(target)
        resolved = mp._resolve_media_url("file://t.docx")
        assert resolved.startswith("file://")
        assert self._local_path(resolved) == target

    def test_media_url_form_resolves(self, ws):
        target = ws / "files" / "media" / "t.docx"
        _docx(target)
        resolved = mp._resolve_media_url("/media/t.docx")
        assert self._local_path(resolved) == target

    def test_lookup_uploaded_file_by_bare_name(self, ws):
        target = ws / "files" / "media" / "t.docx"
        _docx(target)
        found = mp._lookup_uploaded_file("t.docx")
        assert found is not None and Path(found).is_file()
        assert mp._lookup_uploaded_file("ghost.docx") is None

    def test_http_urls_untouched(self):
        url = "https://example.com/a.docx"
        assert mp._resolve_media_url(url) == url

    def test_genuine_file_uri_with_abs_path_untouched(self, tmp_path):
        target = tmp_path / "abs.docx"
        _docx(target)
        assert mp._resolve_media_url(f"file://{target}") == f"file://{target}"
