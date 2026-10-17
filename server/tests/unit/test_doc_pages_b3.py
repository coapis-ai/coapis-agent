"""B3 tests: the paginated document base (parse once, take by need).

Fixtures are generated here, so every number asserted is a number we
measured, not a number we guessed.
"""

from __future__ import annotations

import time
from pathlib import Path

import pytest

from coapis.agents.tools.doc_pages import (
    BLOCK_CHARS,
    XLSX_ROWS_PER_BLOCK,
    cache_key,
    parse_document,
    read_document,
)


# ── fixtures ────────────────────────────────────────────────────────


def _write_pdf(path: Path, pages: int) -> Path:
    import pymupdf

    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        page.insert_text(
            (50, 50),
            f"Page {i + 1} marker-needle-{i + 1} " + ("body text " * 60),
        )
    doc.save(str(path))
    doc.close()
    return path


def _write_xlsx(path: Path, rows: int) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["id", "name", "amount"])
    for i in range(rows):
        ws.append([i, f"row-{i}", i * 3])
    wb.save(str(path))
    return path


def _write_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    for chapter in range(1, 11):
        doc.add_heading(f"第{chapter}章 市场分析", level=1)
        doc.add_paragraph("正文 " * 400)
        doc.add_heading(f"{chapter}.1 细节", level=2)
        doc.add_paragraph(f"needle-{chapter} " + "细节内容 " * 200)
    doc.save(str(path))
    return path


# ── parse once, persist ─────────────────────────────────────────────


def test_pdf_index_is_one_line_per_page(tmp_path):
    pdf = _write_pdf(tmp_path / "big.pdf", 100)
    out = read_document(
        str(pdf),
        "pdf",
        mode="index",
        cache_dir=tmp_path,
    )
    assert out["mode"] == "index"
    assert out["block_count"] == 100
    lines = out["index"].strip().split("\n")
    assert len(lines) == 100
    assert lines[0].startswith("[0] page 第1页")
    # A page block is addressable and never blows the block budget.
    parsed = parse_document(str(pdf), "pdf", cache_dir=tmp_path)["doc"]
    assert all(b.chars <= BLOCK_CHARS for b in parsed.blocks)


def test_pdf_range_reads_exactly_one_page(tmp_path):
    pdf = _write_pdf(tmp_path / "big.pdf", 100)
    index = read_document(str(pdf), "pdf", mode="index", cache_dir=tmp_path)
    page42 = index["blocks"][41]
    out = read_document(
        str(pdf),
        "pdf",
        mode="range",
        start=page42["i"],
        end=page42["i"],
        max_chars=BLOCK_CHARS,
        cache_dir=tmp_path,
    )
    assert "Page 42 marker-needle-42" in out["content"]
    assert out["unit"] == "block"
    assert out["chars"] == page42["chars"]
    assert out["char_span"] == [page42["start"], page42["end"]]
    assert out["blocks_returned"] == [page42["i"]]


def test_block_range_walks_the_document_without_mixing_units(tmp_path):
    """Block numbers page through; character offsets are a separate door."""
    pdf = _write_pdf(tmp_path / "big.pdf", 100)
    index = read_document(str(pdf), "pdf", mode="index", cache_dir=tmp_path)
    out = read_document(
        str(pdf),
        "pdf",
        mode="range",
        start=0,
        end=99,
        max_chars=BLOCK_CHARS,
        cache_dir=tmp_path,
    )
    assert out["unit"] == "block"
    assert out["content"].startswith("Page 1 ")
    # The cursor continues in blocks, never in characters.
    assert out["cursor"]["next_block"] == len(out["blocks_returned"])
    assert "start=" in out["cursor"]["hint"]
    assert "start_char" not in out["cursor"]["hint"]
    # Character offsets stay available under their own names.
    chars = read_document(
        str(pdf),
        "pdf",
        mode="range",
        start_char=0,
        end_char=200,
        max_chars=BLOCK_CHARS,
        cache_dir=tmp_path,
    )
    assert chars["unit"] == "char"
    assert chars["chars"] == 200
    assert "start_char=" in chars["cursor"]["hint"]


def test_second_read_is_a_cache_hit_with_zero_parsing(tmp_path):
    pdf = _write_pdf(tmp_path / "big.pdf", 30)
    first = parse_document(str(pdf), "pdf", cache_dir=tmp_path)
    assert "doc" in first
    assert first["doc"].cached is True  # written on the first parse

    files = list(tmp_path.glob("*.md"))
    assert len(files) == 1, "one canonical markdown per document"

    second = parse_document(str(pdf), "pdf", cache_dir=tmp_path)
    assert second["doc"].cached is True
    assert second["doc"].key == first["doc"].key
    assert second["doc"].text == first["doc"].text
    assert [b.end for b in second["doc"].blocks] == [
        b.end for b in first["doc"].blocks
    ]


def test_cache_key_follows_mtime(tmp_path):
    doc = _write_docx(tmp_path / "a.docx")
    key_before = cache_key(str(doc), "docx")
    time.sleep(0.01)
    os_touch(doc)
    assert cache_key(str(doc), "docx") != key_before, (
        "a re-uploaded file must not read from the stale cache"
    )


def os_touch(path: Path) -> None:
    now = time.time()
    import os

    os.utime(path, (now, now))


# ── per-format block shapes ─────────────────────────────────────────


def test_xlsx_sheet_blocks_carry_row_ranges(tmp_path):
    xlsx = _write_xlsx(tmp_path / "rows.xlsx", 5000)
    out = read_document(
        str(xlsx),
        "xlsx",
        mode="index",
        cache_dir=tmp_path,
    )
    assert out["block_count"] > 1, "5000 rows must not be one block"
    rows = [b for b in out["index"].split("\n") if b.strip()]
    assert "Sheet: Ledger" in rows[0]
    parsed = parse_document(str(xlsx), "xlsx", cache_dir=tmp_path)["doc"]
    for block in parsed.blocks:
        assert block.meta.get("sheet") == "Ledger"
        assert block.meta.get("rows", 0) <= XLSX_ROWS_PER_BLOCK
        assert block.chars <= BLOCK_CHARS + 2000
    assert sum(b.meta.get("rows", 0) for b in parsed.blocks) >= 5000


def test_docx_blocks_are_heading_sections(tmp_path):
    doc = _write_docx(tmp_path / "doc.docx")
    parsed = parse_document(str(doc), "docx", cache_dir=tmp_path)["doc"]
    titles = [b.title for b in parsed.blocks]
    assert any("第1章 市场分析" in t for t in titles)
    assert any("needle-7" in parsed.text[b.start : b.end] for b in parsed.blocks)
    # Level-1 and level-2 headings both open a section.
    assert any("1.1 细节" in t for t in titles)


def test_text_without_headings_splits_by_size(tmp_path):
    raw = tmp_path / "plain.txt"
    raw.write_text("行内容 " * 4000, encoding="utf-8")
    parsed = parse_document(str(raw), "txt", cache_dir=tmp_path)["doc"]
    assert len(parsed.blocks) > 1
    assert all(b.chars <= BLOCK_CHARS for b in parsed.blocks)
    assert parsed.text.startswith("行内容")


# ── the four modes ──────────────────────────────────────────────────


def test_range_is_capped_by_the_budget_and_says_so(tmp_path):
    doc = _write_docx(tmp_path / "doc.docx")
    index = read_document(str(doc), "docx", mode="index", cache_dir=tmp_path)
    total = index["total_chars"]
    out = read_document(
        str(doc),
        "docx",
        mode="range",
        start_char=0,
        end_char=total,
        max_chars=5000,
        cache_dir=tmp_path,
    )
    assert out["chars"] == 5000
    assert out["requested_chars"] == total
    assert out["over_budget"] is True
    assert out["cursor"]["read_to"] == 5000
    assert out["cursor"]["next_start"] == 5000
    assert out["cursor"]["remaining_chars"] == total - 5000
    assert "start_char=5000" in out["cursor"]["hint"]


def test_ole2_file_named_docx_is_named_not_misread(tmp_path):
    """A Word 97-2003 file wearing a .docx name must say what it is."""
    fake = tmp_path / "legacy.docx"
    fake.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    out = read_document(str(fake), "docx", mode="index", cache_dir=tmp_path)
    assert "OLE2" in out["error"]
    assert out["format_detected"] == "ole2"


def test_range_past_the_end_is_an_explicit_answer(tmp_path):
    raw = tmp_path / "plain.txt"
    raw.write_text("短文档", encoding="utf-8")
    out = read_document(
        str(raw),
        "txt",
        mode="range",
        start=999,
        cache_dir=tmp_path,
    )
    assert "超出块数" in out["error"]


def test_search_returns_context_with_addresses(tmp_path):
    pdf = _write_pdf(tmp_path / "big.pdf", 60)
    out = read_document(
        str(pdf),
        "pdf",
        mode="search",
        keywords="marker-needle-37",
        max_chars=20000,
        cache_dir=tmp_path,
    )
    assert out["mode"] == "search"
    assert out["hits"] >= 1
    hit = out["results"][0]
    assert "marker-needle-37" in hit["text"]
    assert hit["block"] == 36, "the hit must point at page 37's block"
    assert hit["end"] - hit["start"] <= 400 + len("marker-needle-37")


def test_search_with_no_hit_says_so(tmp_path):
    raw = tmp_path / "plain.txt"
    raw.write_text("这里没有任何目标", encoding="utf-8")
    out = read_document(
        str(raw),
        "txt",
        mode="search",
        keywords="不存在",
        cache_dir=tmp_path,
    )
    assert out["hits"] == 0
    assert "未命中" in out["hint"]


def test_full_mode_caps_and_announces_the_cap(tmp_path):
    doc = _write_docx(tmp_path / "doc.docx")
    out = read_document(
        str(doc),
        "docx",
        mode="full",
        max_chars=8000,
        cache_dir=tmp_path,
    )
    assert out["chars"] == 8000
    assert out["truncated"] is True
    assert "预算上限" in out["warning"]
    assert out["cursor"]["next_start"] == 8000


def test_unknown_mode_is_rejected(tmp_path):
    raw = tmp_path / "plain.txt"
    raw.write_text("x", encoding="utf-8")
    out = read_document(str(raw), "txt", mode="teleport", cache_dir=tmp_path)
    assert "无效" in out["error"]


def test_missing_dependency_names_the_install_route(tmp_path):
    fake = tmp_path / "x.docx"
    fake.write_bytes(b"not a docx")
    out = read_document(str(fake), "docx", mode="index", cache_dir=tmp_path)
    assert "error" in out
    assert "install_dependency" in out["error"] or "解析失败" in out["error"]
