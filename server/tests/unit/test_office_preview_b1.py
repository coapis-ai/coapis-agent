# -*- coding: utf-8 -*-
"""B1 回归：预览抽屉必须能渲染 office 文档与 PDF。

背景（2026-10-22 企业环境实测）：预览抽屉只认文本/图片，`docx/xlsx/pptx/pdf`
一律掉进 `"other"` → 用户点预览吃到「不支持在线预览」。企业场景的交付物恰恰
以 docx 为主，前端白名单里也列了这些扩展名（前端认为可预览、后端说不能），
属于「前后端一套做一套」。

锁死四件事：
 1. docx/xlsx/pptx → `type=="text"` 且 content 有正文（走 markdown 渲染分支）
 2. pdf → `type=="pdf"`（抽屉内联渲染）
 3. 解析失败/损坏文件 → 降级为 `"other"`，**不得抛 500**
 4. 旧二进制格式 .doc/.xls/.ppt → 降级 `"other"`（走下载），不炸
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from coapis.app.services.file_service import LocalFileBackend


# ── 造真实 office 文件（镜像里 python-docx/openpyxl/python-pptx 齐备）──


def _make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("入市方案", level=1)
    doc.add_paragraph("这是正文第一段。")
    doc.add_paragraph("第二段：结论先行。")
    doc.save(str(path))


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "分数"])
    ws.append(["张三", 92])
    ws.append(["李四", 87])
    wb.save(str(path))


def _make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "季度复盘"
    box = slide.shapes.add_textbox(
        Inches(1), Inches(2), Inches(6), Inches(2)
    )
    box.text_frame.text = "三条要点：增长、留存、成本。"
    prs.save(str(path))


@pytest.fixture()
def backend(tmp_path):
    root = tmp_path / "workspaces" / "alice" / "files"
    root.mkdir(parents=True)
    _make_docx(root / "plan.docx")
    _make_xlsx(root / "score.xlsx")
    _make_pptx(root / "deck.pptx")
    (root / "doc.pdf").write_bytes(b"%PDF-1.4\n%\xff\xff\xff\xff\n1 0 obj\nendobj\n")
    (root / "legacy.doc").write_bytes(b"\xd0\xcf\x11\xe0fake-msword")
    (root / "broken.docx").write_bytes(b"PK not really a zip")
    return LocalFileBackend(tmp_path / "workspaces")


def _preview(backend, name):
    return asyncio.run(backend.preview_file("alice", name, "files"))


# ── office 必须出正文 ─────────────────────────────────────────────


def test_docx_preview_returns_text(backend):
    res = _preview(backend, "plan.docx")
    assert res["type"] == "text"
    assert "入市方案" in res["content"]
    assert "正文第一段" in res["content"]
    assert res["converted_from"] == ".docx"


def test_xlsx_preview_returns_text(backend):
    res = _preview(backend, "score.xlsx")
    assert res["type"] == "text"
    assert "姓名" in res["content"]
    assert "张三" in res["content"]


def test_pptx_preview_returns_text(backend):
    res = _preview(backend, "deck.pptx")
    assert res["type"] == "text"
    assert "季度复盘" in res["content"]


def test_pdf_preview_returns_pdf_type(backend):
    res = _preview(backend, "doc.pdf")
    assert res["type"] == "pdf"
    assert res["name"] == "doc.pdf"


# ── 失败必须降级，不得抛 500 ──────────────────────────────────────


def test_broken_docx_degrades_to_other(backend):
    """损坏文件解析失败 → 'other'（前端显示下载），不得抛异常。

    容器闸门：损坏 docx 不是合法 zip。若放它进 markitdown 兜底，兜底会把
    **原始字节当文本**返回（实测预览正文 = "PK not really a zip"），用户
    点预览看到的是一串乱码。因此预览正文必须是真正文，乱码一律降级下载。
    """
    res = _preview(backend, "broken.docx")
    assert res["type"] == "other"
    assert "content" not in res


def test_legacy_doc_degrades_to_other(backend):
    """旧二进制格式无解析器 → 'other'，走下载。"""
    res = _preview(backend, "legacy.doc")
    assert res["type"] == "other"


def test_office_preview_timeout_guard(backend, monkeypatch):
    """超时护栏：解析卡死时降级为 other，不能把请求吊死。"""
    from coapis.app.services import file_service as fs

    monkeypatch.setattr(fs, "_OFFICE_PREVIEW_TIMEOUT_S", 0.01)

    async def _slow(*a, **kw):
        await asyncio.sleep(5)

    import coapis.agents.tools.doc_reader as dr

    monkeypatch.setattr(dr, "_read_docx", _slow)
    res = _preview(backend, "plan.docx")
    assert res["type"] == "other"
